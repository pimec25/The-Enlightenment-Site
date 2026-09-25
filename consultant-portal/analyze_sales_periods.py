"""Compare sales order detail by period and find gross-margin and Pareto signals."""

import argparse
import json
from collections import defaultdict
from datetime import datetime, date as date_type
from pathlib import Path

from openpyxl import load_workbook


def pct(n, d):
    return n / d if d else None


def metrics(rows):
    sales = sum(r["sales"] for r in rows)
    cost = sum(r["cost"] for r in rows)
    gp = sales - cost
    return {"sales": sales, "cost": cost, "gross_profit": gp,
            "gross_margin_pct": 100 * pct(gp, sales) if sales else None,
            "row_count": len(rows)}


def pareto(rows, field, metric):
    grouped = defaultdict(list)
    for row in rows:
        grouped[row[field]].append(row)
    ranked = sorted(((key, metrics(values)) for key, values in grouped.items()),
                    key=lambda x: -x[1][metric])
    target = 0.8 * sum(max(0, values[metric]) for _, values in ranked)
    cumulative = 0.0
    selected = []
    for key, values in ranked:
        if cumulative >= target:
            break
        cumulative += max(0, values[metric])
        selected.append({field: key, **values})
    return {"total_groups": len(ranked), "groups_to_80": len(selected),
            "pct_groups_to_80": 100 * pct(len(selected), len(ranked)) if ranked else None,
            "amount_covered": cumulative, "top_10": selected[:10]}


def analyze(path, as_of=None):
    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book.active
    iterator = sheet.iter_rows(values_only=True)
    headers = next(iterator)
    needed = ["Request Date", "Extended Price", "Extended Cost", "Parent Number",
              "2nd Item Number", "Value Stream Product Family", "Order Number", "Line Number"]
    missing = [name for name in needed if name not in headers]
    if missing:
        raise ValueError(f"Missing columns: {missing}")
    ix = {name: headers.index(name) for name in needed}
    rows = []
    for excel_row, values in enumerate(iterator, 2):
        date = values[ix["Request Date"]]
        if not isinstance(date, datetime):
            continue
        rows.append({"excel_row": excel_row, "date": date, "year": date.year,
                     "sales": float(values[ix["Extended Price"]] or 0),
                     "cost": float(values[ix["Extended Cost"]] or 0),
                     "customer": str(values[ix["Parent Number"]]).strip(),
                     "item": str(values[ix["2nd Item Number"]]).strip(),
                     "value_stream": str(values[ix["Value Stream Product Family"]]).strip(),
                     "order": str(values[ix["Order Number"]]).strip(),
                     "line": str(values[ix["Line Number"]]).strip()})
    if not rows:
        raise ValueError("No rows with valid Request Date")
    as_of = as_of or date_type.today()
    current_year = as_of.year
    cutoff = (as_of.month, as_of.day)
    current = [r for r in rows if r["year"] == current_year and (r["date"].month, r["date"].day) <= cutoff]
    prior = [r for r in rows if r["year"] == current_year - 1 and (r["date"].month, r["date"].day) <= cutoff]
    cur, prv = metrics(current), metrics(prior)
    loss_rows = sorted((dict(r) for r in current if r["cost"] > r["sales"]),
                       key=lambda r: -(r["cost"] - r["sales"]))
    for row in loss_rows:
        row["gross_loss"] = row["cost"] - row["sales"]
        if row["sales"] == 0 and row["cost"] > 0:
            row["review_class"] = "Zero-price shipment / positive cost"
        elif row["sales"] < 0 and row["cost"] < 0:
            row["review_class"] = "Return or credit with margin impact"
        elif row["sales"] < 0 and row["cost"] >= 0:
            row["review_class"] = "Credit with unreversed cost"
        else:
            row["review_class"] = "Positive-price sale below cost"
        row["date"] = row["date"].date().isoformat()
    zero_sales = [r["excel_row"] for r in current if r["sales"] == 0 and r["cost"] != 0]
    stream = defaultdict(list)
    prior_stream = defaultdict(list)
    for r in current:
        stream[r["value_stream"]].append(r)
    for r in prior:
        prior_stream[r["value_stream"]].append(r)
    stream_comparison = {}
    for key in sorted(set(stream) | set(prior_stream)):
        cm, pm = metrics(stream[key]), metrics(prior_stream[key])
        stream_comparison[key] = {"current": cm, "prior": pm,
                                  "sales_change": cm["sales"] - pm["sales"],
                                  "gross_profit_change": cm["gross_profit"] - pm["gross_profit"],
                                  "gross_margin_point_change": ((cm["gross_margin_pct"] or 0) - (pm["gross_margin_pct"] or 0))}
    result = {"source": str(Path(path).resolve()), "sheet": sheet.title,
            "date_field": "Request Date", "as_of": as_of.isoformat(),
            "future_dated_rows_excluded": sum(r["date"].date() > as_of for r in rows),
            "current_period": {"label": f"{current_year}-01-01 through {as_of.isoformat()}", **cur},
            "prior_comparable_period": {"label": f"{current_year-1}-01-01 through {current_year-1}-{cutoff[0]:02d}-{cutoff[1]:02d}", **prv},
            "change": {"sales": cur["sales"] - prv["sales"],
                       "sales_pct": 100 * pct(cur["sales"]-prv["sales"], prv["sales"]) if prv["sales"] else None,
                       "gross_profit": cur["gross_profit"] - prv["gross_profit"],
                       "gross_profit_pct": 100 * pct(cur["gross_profit"]-prv["gross_profit"], prv["gross_profit"]) if prv["gross_profit"] else None,
                       "gross_margin_points": cur["gross_margin_pct"] - prv["gross_margin_pct"]},
            "pareto": {field: {metric: pareto(current, field, metric) for metric in ["sales", "gross_profit"]}
                       for field in ["customer", "item"]},
            "current_value_streams": {key: metrics(values) for key, values in stream.items()},
            "value_stream_comparison": stream_comparison,
            "current_below_cost_rows": len(loss_rows),
            "current_below_cost_gross_loss": sum(r["gross_loss"] for r in loss_rows),
            "top_20_current_below_cost_rows": loss_rows[:20],
            "current_below_cost_detail": loss_rows,
            "current_zero_sales_nonzero_cost_count": len(zero_sales),
            "current_zero_sales_nonzero_cost_rows": zero_sales[:100],
            "limitations": ["This is sales-order detail, not a complete P&L; operating expenses and EBITDA are absent.",
                            "Request Date is used because it is consistently populated; it may not equal revenue-recognition or invoice date.",
                            "Below-cost and margin movements are review signals, not confirmed errors or recoverable savings."]}
    book.close()
    return result


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--as-of", type=date_type.fromisoformat, help="Comparison cutoff in YYYY-MM-DD format")
    args = parser.parse_args()
    result = analyze(args.workbook, args.as_of)
    text = json.dumps(result, indent=2)
    if args.output:
        args.output.write_text(text + "\n", encoding="utf-8")
    else:
        print(text)

